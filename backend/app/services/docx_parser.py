import docx
import io
import re
from backend.app.services.pdf_parser import clean_extracted_text

def extract_text_from_docx(file_bytes: bytes) -> dict:
    """
    Extracts text from DOCX file bytes using python-docx.
    """
    try:
        doc_stream = io.BytesIO(file_bytes)
        doc = docx.Document(doc_stream)
        
        full_text_paragraphs = []
        for paragraph in doc.paragraphs:
            if paragraph.text:
                full_text_paragraphs.append(paragraph.text)
                
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text_paragraphs.append(" | ".join(row_text))
                    
        raw_text = "\n".join(full_text_paragraphs)
        cleaned_text = clean_extracted_text(raw_text)
        
        if not cleaned_text:
            return {
                "extracted_text": "",
                "warning": "DOCX file contains no readable text content.",
                "char_count": 0,
                "word_count": 0
            }
            
        return {
            "extracted_text": cleaned_text,
            "warning": None,
            "char_count": len(cleaned_text),
            "word_count": len(cleaned_text.split())
        }
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")
